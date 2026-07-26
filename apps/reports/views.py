from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.core.exceptions import ValidationError
from apps.accounts.decorators import role_required
from apps.accounts.models import LecturerProfile, MentorProfile, DepartmentProfile
from .models import (
    InternshipAssignment, WeeklyReport, FinalReport, LecturerFeedback,
    InternshipTerm, ReportComment,
)
from .forms import (
    WeeklyReportForm,
    FinalReportForm,
    LecturerFeedbackForm,
    MentorGradeForm,
    LecturerGradeForm,
    FinalReportMentorGradeForm,
    FinalReportLecturerGradeForm,
    ReportCommentForm,
)
from django.utils import timezone
from apps.reports.services.submission_service import submit_weekly, submit_final, unlock_report
from apps.reports.services.grading_service import (
    grade_weekly_report_by_mentor,
    grade_weekly_report_by_lecturer,
    grade_final_report_by_mentor,
    grade_final_report_by_lecturer,
    calculate_assignment_score,
)
from apps.reports.services.statistics_service import build_statistics_context
import csv
from django.http import HttpResponse
import openpyxl
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.styles import Font, Alignment

# =================== HELPERS ===================

def get_student_assignment(request):
    """Lấy assignment của student đang đăng nhập cho kỳ hiện tại."""
    from django.db.models import Q
    base_qs = InternshipAssignment.objects.filter(
        Q(application__student=request.user.student_profile, is_active=True) |
        Q(external_request__student=request.user.student_profile, is_active=True)
    )
    active_term = InternshipTerm.get_active()
    if active_term:
        term_match = base_qs.filter(term=active_term).first()
        if term_match:
            return term_match
    return base_qs.first()


def get_selected_term(request):
    """Helper lấy term đang chọn từ query param hoặc fallback active term."""
    all_terms = InternshipTerm.objects.all().order_by('-year')
    active_term = InternshipTerm.get_active()
    term_id = request.GET.get('period') or request.GET.get('term')
    selected_term = None
    if term_id:
        try:
            selected_term = InternshipTerm.objects.get(pk=term_id)
        except InternshipTerm.DoesNotExist:
            selected_term = active_term
    else:
        selected_term = active_term
    return selected_term, all_terms, active_term


# =================== STUDENT VIEWS ===================

@role_required(['student'])
def student_dashboard(request):
    assignment = get_student_assignment(request)
    if not assignment:
        return render(request, 'reports/student_no_assignment.html')

    weekly_reports = assignment.weekly_reports.all()
    submitted_count = weekly_reports.filter(is_submitted=True).count()
    has_final = hasattr(assignment, 'final_report')

    return render(request, 'reports/student_dashboard.html', {
        'assignment': assignment,
        'weekly_count': submitted_count,
        'has_final': has_final,
    })


@role_required(['student'])
def weekly_report_list(request):
    assignment = get_student_assignment(request)
    if not assignment:
        return render(request, 'reports/student_no_assignment.html')

    reports = assignment.weekly_reports.prefetch_related('grades').all()
    for report in reports:
        report.mentor_grade = report.grades.filter(grader_type='mentor').first()
        report.lecturer_grade = report.grades.filter(grader_type='lecturer').first()

    return render(request, 'reports/weekly_report_list.html', {'reports': reports, 'assignment': assignment})


@role_required(['student'])
def weekly_report_submit(request, week_number):
    """SV chọn tuần → điền nội dung → submit. Report đã tạo sẵn."""
    assignment = get_student_assignment(request)
    if not assignment:
        return render(request, 'reports/student_no_assignment.html')

    report = get_object_or_404(WeeklyReport, assignment=assignment, week_number=week_number)

    if report.is_locked:
        messages.error(request, f'Báo cáo Tuần {week_number} đã bị khóa. Không thể nộp lại.')
        return redirect('reports:weekly_report_list')

    if request.method == 'POST':
        form = WeeklyReportForm(request.POST, request.FILES, instance=report)
        if form.is_valid():
            form.save()
            try:
                submit_weekly(report, request.user)
                messages.success(request, f'Đã nộp báo cáo tuần {week_number} thành công!')
            except ValidationError as e:
                messages.error(request, str(e.message))
            return redirect('reports:weekly_report_list')
    else:
        form = WeeklyReportForm(instance=report)

    return render(request, 'reports/weekly_report_form.html', {
        'form': form,
        'assignment': assignment,
        'report': report,
        'week_number': week_number,
    })


@role_required(['student'])
def final_report_detail(request):
    assignment = get_student_assignment(request)
    if not assignment:
        return render(request, 'reports/student_no_assignment.html')

    try:
        report = assignment.final_report
    except FinalReport.DoesNotExist:
        messages.error(request, 'Chưa có báo cáo tổng kết. Vui lòng liên hệ quản trị viên.')
        return redirect('reports:student_dashboard')

    if report.is_submitted:
        mentor_grade = report.grades.filter(grader_type='mentor').first()
        lecturer_grade = report.grades.filter(grader_type='lecturer').first()
        score_info = calculate_assignment_score(assignment)
        return render(request, 'reports/final_report_detail.html', {
            'report': report,
            'assignment': assignment,
            'mentor_grade': mentor_grade,
            'lecturer_grade': lecturer_grade,
            'total_score': score_info.get('total_score'),
            'is_passed': score_info.get('is_passed'),
        })

    if request.method == 'POST':
        form = FinalReportForm(request.POST, request.FILES, instance=report)
        if form.is_valid():
            form.save()
            try:
                submit_final(report, request.user)
                messages.success(request, 'Đã hoàn tất nộp báo cáo tổng kết đợt thực tập!')
            except ValidationError as e:
                messages.error(request, str(e.message))
            return redirect('reports:final_report_detail')
    else:
        form = FinalReportForm(instance=report)

    return render(request, 'reports/final_report_form.html', {'form': form, 'assignment': assignment, 'report': report})


# =================== MENTOR VIEWS ===================

@role_required(['mentor'])
def mentor_dashboard(request):
    mentor_profile, _ = MentorProfile.objects.get_or_create(user=request.user)
    selected_term, all_terms, active_term = get_selected_term(request)
    status_filter = request.GET.get('status', 'all')

    assignments = InternshipAssignment.objects.filter(
        mentor_group__mentor=mentor_profile
    ).select_related(
        'application__student__user', 'application__post__department__user',
        'external_request__student__user', 'term', 'mentor_group',
    )
    if selected_term:
        assignments = assignments.filter(term=selected_term)

    weekly_reports = WeeklyReport.objects.filter(
        assignment__mentor_group__mentor=mentor_profile, is_submitted=True
    ).select_related(
        'assignment__application__student__user', 'assignment__external_request__student__user',
        'assignment__application__post__department', 'assignment__term',
    ).prefetch_related('grades').order_by('-submitted_at', '-week_number')
    if selected_term:
        weekly_reports = weekly_reports.filter(assignment__term=selected_term)

    if status_filter == 'pending':
        weekly_reports = weekly_reports.exclude(grades__grader_type='mentor')
    elif status_filter == 'graded':
        weekly_reports = weekly_reports.filter(grades__grader_type='mentor')

    for report in weekly_reports:
        report.mentor_grade = report.grades.filter(grader_type='mentor').first()

    final_reports = FinalReport.objects.filter(
        assignment__mentor_group__mentor=mentor_profile, is_submitted=True
    ).select_related(
        'assignment__application__student__user', 'assignment__external_request__student__user',
        'assignment__application__post__department', 'assignment__term',
    ).prefetch_related('grades').order_by('-submitted_at')
    if selected_term:
        final_reports = final_reports.filter(assignment__term=selected_term)

    if status_filter == 'pending':
        final_reports = final_reports.exclude(grades__grader_type='mentor')
    elif status_filter == 'graded':
        final_reports = final_reports.filter(grades__grader_type='mentor')

    for report in final_reports:
        report.mentor_grade = report.grades.filter(grader_type='mentor').first()
        report.lecturer_grade = report.grades.filter(grader_type='lecturer').first()

    return render(request, 'reports/mentor_dashboard.html', {
        'assignments': assignments,
        'weekly_reports': weekly_reports,
        'final_reports': final_reports,
        'all_periods': all_terms,
        'selected_period': selected_term,
        'active_period': active_term,
        'status_filter': status_filter,
        'metrics': {
            'assignment_count': assignments.count(),
            'total_weekly': weekly_reports.count(),
            'pending_weekly': weekly_reports.exclude(grades__grader_type='mentor').count() if status_filter == 'all' else 0,
            'late_weekly': weekly_reports.filter(is_late=True).count(),
        }
    })


@role_required(['mentor', 'lecturer'])
def weekly_report_detail(request, report_id):
    selected_term, all_terms, active_term = get_selected_term(request)
    is_mentor = request.user.role == 'mentor'

    if is_mentor:
        mentor_profile, _ = MentorProfile.objects.get_or_create(user=request.user)
        report = get_object_or_404(
            WeeklyReport.objects.select_related(
                'assignment__mentor_group__mentor', 'assignment__lecturer', 'assignment__term',
                'assignment__application__student__user', 'assignment__external_request__student__user',
                'assignment__application__post__department',
            ).prefetch_related('grades'),
            id=report_id, assignment__mentor_group__mentor=mentor_profile,
        )
    else:
        lecturer_profile, _ = LecturerProfile.objects.get_or_create(user=request.user)
        report = get_object_or_404(
            WeeklyReport.objects.select_related(
                'assignment__mentor_group__mentor', 'assignment__lecturer', 'assignment__term',
                'assignment__application__student__user', 'assignment__external_request__student__user',
                'assignment__application__post__department',
            ).prefetch_related('grades'),
            id=report_id, assignment__lecturer=lecturer_profile,
        )

    assignment = report.assignment
    if selected_term and assignment.term != selected_term:
        return redirect('reports:mentor_dashboard' if is_mentor else 'reports:lecturer_dashboard')

    mentor_grade = report.grades.filter(grader_type='mentor').first()
    lecturer_grade = report.grades.filter(grader_type='lecturer').first()

    if request.method == 'POST':
        form = MentorGradeForm(request.POST) if is_mentor else LecturerGradeForm(request.POST)
        if form.is_valid():
            score = form.cleaned_data['score']
            comment = form.cleaned_data['comment']
            if is_mentor:
                grade_weekly_report_by_mentor(report, mentor_profile, score, comment)
                messages.success(request, f'Đã lưu điểm Mentor cho Báo cáo Tuần {report.week_number}.')
            else:
                grade_weekly_report_by_lecturer(report, lecturer_profile, score, comment)
                messages.success(request, f'Đã lưu điểm Giảng viên cho Báo cáo Tuần {report.week_number}.')
            return redirect('reports:weekly_report_detail', report_id=report_id)
        grade_form = form
    else:
        initial_data = {}
        if is_mentor and mentor_grade:
            initial_data = {'score': mentor_grade.score, 'comment': mentor_grade.comment}
        if not is_mentor and lecturer_grade:
            initial_data = {'score': lecturer_grade.score, 'comment': lecturer_grade.comment}
        grade_form = MentorGradeForm(initial=initial_data) if is_mentor else LecturerGradeForm(initial=initial_data)

    return render(request, 'reports/weekly_report_detail.html', {
        'report': report, 'assignment': assignment,
        'mentor_grade': mentor_grade, 'lecturer_grade': lecturer_grade,
        'grade_form': grade_form, 'is_mentor': is_mentor,
        'all_periods': all_terms, 'selected_period': selected_term, 'active_period': active_term,
    })


@role_required(['mentor', 'lecturer'])
def final_report_grade_detail(request, report_id):
    is_mentor = request.user.role == 'mentor'
    selected_term, all_terms, active_term = get_selected_term(request)

    if is_mentor:
        mentor_profile, _ = MentorProfile.objects.get_or_create(user=request.user)
        report = get_object_or_404(
            FinalReport.objects.select_related(
                'assignment__mentor_group__mentor', 'assignment__lecturer', 'assignment__term',
                'assignment__application__student__user', 'assignment__external_request__student__user',
                'assignment__application__post__department',
            ).prefetch_related('grades'),
            id=report_id, assignment__mentor_group__mentor=mentor_profile,
        )
    else:
        lecturer_profile, _ = LecturerProfile.objects.get_or_create(user=request.user)
        report = get_object_or_404(
            FinalReport.objects.select_related(
                'assignment__mentor_group__mentor', 'assignment__lecturer', 'assignment__term',
                'assignment__application__student__user', 'assignment__external_request__student__user',
                'assignment__application__post__department',
            ).prefetch_related('grades'),
            id=report_id, assignment__lecturer=lecturer_profile,
        )

    assignment = report.assignment
    if selected_term and assignment.term != selected_term:
        return redirect('reports:mentor_dashboard' if is_mentor else 'reports:lecturer_dashboard')

    mentor_grade = report.grades.filter(grader_type='mentor').first()
    lecturer_grade = report.grades.filter(grader_type='lecturer').first()

    if request.method == 'POST':
        form = FinalReportMentorGradeForm(request.POST) if is_mentor else FinalReportLecturerGradeForm(request.POST)
        if form.is_valid():
            score = form.cleaned_data['score']
            comment = form.cleaned_data['comment']
            if is_mentor:
                grade_final_report_by_mentor(report, mentor_profile, score, comment)
                messages.success(request, 'Đã lưu điểm Mentor cho Báo cáo Tổng kết.')
            else:
                grade_final_report_by_lecturer(report, lecturer_profile, score, comment)
                messages.success(request, 'Đã lưu điểm Giảng viên cho Báo cáo Tổng kết.')
            return redirect('reports:final_report_grade_detail', report_id=report_id)
        grade_form = form
    else:
        initial_data = {}
        if is_mentor and mentor_grade:
            initial_data = {'score': mentor_grade.score, 'comment': mentor_grade.comment}
        if not is_mentor and lecturer_grade:
            initial_data = {'score': lecturer_grade.score, 'comment': lecturer_grade.comment}
        grade_form = FinalReportMentorGradeForm(initial=initial_data) if is_mentor else FinalReportLecturerGradeForm(initial=initial_data)

    return render(request, 'reports/final_report_grade_detail.html', {
        'report': report, 'assignment': assignment,
        'mentor_grade': mentor_grade, 'lecturer_grade': lecturer_grade,
        'grade_form': grade_form, 'is_mentor': is_mentor,
        'all_periods': all_terms, 'selected_period': selected_term, 'active_period': active_term,
    })


# =================== LECTURER VIEWS ===================

@role_required(['lecturer'])
def lecturer_dashboard(request):
    from django.db.models import Count, Q

    lecturer_profile, _ = LecturerProfile.objects.get_or_create(user=request.user)
    selected_term, all_terms, active_term = get_selected_term(request)

    assignments = InternshipAssignment.objects.filter(
        lecturer=lecturer_profile
    ).select_related(
        'application__student__user', 'application__post__department__user',
        'external_request__student__user', 'term',
    )
    if selected_term:
        assignments = assignments.filter(term=selected_term)

    search_q = request.GET.get('q', '').strip()
    if search_q:
        assignments = assignments.filter(
            Q(application__student__user__first_name__icontains=search_q) |
            Q(application__student__user__last_name__icontains=search_q) |
            Q(application__student__user__username__icontains=search_q) |
            Q(external_request__student__user__first_name__icontains=search_q) |
            Q(external_request__student__user__last_name__icontains=search_q) |
            Q(external_request__student__user__username__icontains=search_q)
        )

    search_company = request.GET.get('company', '').strip()
    if search_company:
        assignments = assignments.filter(
            Q(application__post__department__organization_name__icontains=search_company) |
            Q(application__post__department__user__username__icontains=search_company) |
            Q(external_request__company_name__icontains=search_company)
        )

    assignments = assignments.order_by('-start_date')
    total_students = assignments.count()

    all_weekly_reports = WeeklyReport.objects.filter(assignment__lecturer=lecturer_profile, is_submitted=True)
    if selected_term:
        all_weekly_reports = all_weekly_reports.filter(assignment__term=selected_term)
    total_weekly = all_weekly_reports.count()
    weekly_waiting = all_weekly_reports.exclude(grades__grader_type='lecturer').count()

    selected_week = None
    week_number = request.GET.get('week')
    if week_number:
        try:
            selected_week = int(week_number)
        except ValueError:
            pass

    status_filter = request.GET.get('status', 'all')
    weekly_reports = WeeklyReport.objects.filter(assignment__lecturer=lecturer_profile, is_submitted=True)
    if selected_term:
        weekly_reports = weekly_reports.filter(assignment__term=selected_term)
    if selected_week:
        weekly_reports = weekly_reports.filter(week_number=selected_week)
    if status_filter == 'pending':
        weekly_reports = weekly_reports.exclude(grades__grader_type='lecturer')
    elif status_filter == 'graded':
        weekly_reports = weekly_reports.filter(grades__grader_type='lecturer')

    weekly_reports = weekly_reports.select_related(
        'assignment__application__student__user', 'assignment__external_request__student__user',
        'assignment__application__post__department', 'assignment__mentor_group', 'assignment__term',
    ).prefetch_related('grades').order_by('-submitted_at', '-week_number').distinct()

    for report in weekly_reports:
        report.mentor_grade = report.grades.filter(grader_type='mentor').first()
        report.lecturer_grade = report.grades.filter(grader_type='lecturer').first()

    final_reports = FinalReport.objects.filter(assignment__lecturer=lecturer_profile, is_submitted=True)
    if selected_term:
        final_reports = final_reports.filter(assignment__term=selected_term)
    if search_company:
        final_reports = final_reports.filter(
            Q(assignment__application__post__department__organization_name__icontains=search_company) |
            Q(assignment__application__post__department__user__username__icontains=search_company) |
            Q(assignment__external_request__company_name__icontains=search_company)
        )
    final_reports = final_reports.select_related(
        'assignment__application__student__user', 'assignment__external_request__student__user',
        'assignment__application__post__department', 'assignment__mentor_group__mentor',
        'assignment__lecturer', 'assignment__term',
    ).prefetch_related('grades').order_by('-submitted_at')

    for report in final_reports:
        report.mentor_grade = report.grades.filter(grader_type='mentor').first()
        report.lecturer_grade = report.grades.filter(grader_type='lecturer').first()

    total_final = final_reports.count()
    final_waiting = final_reports.exclude(grades__grader_type='lecturer').count()

    students_no_reports = assignments.annotate(
        num_submitted=Count('weekly_reports', filter=Q(weekly_reports__is_submitted=True))
    ).filter(num_submitted=0).count()

    context = {
        'assignments': assignments,
        'all_periods': all_terms, 'selected_period': selected_term, 'active_period': active_term,
        'search_q': search_q, 'search_company': search_company,
        'selected_week': selected_week, 'status_filter': status_filter,
        'weekly_reports': weekly_reports, 'final_reports': final_reports,
        'metrics': {
            'total_students': total_students, 'total_weekly': total_weekly,
            'weekly_waiting': weekly_waiting, 'total_final': total_final,
            'final_waiting': final_waiting, 'students_no_reports': students_no_reports,
            'total_waiting': weekly_waiting + final_waiting,
        }
    }
    return render(request, 'reports/lecturer_dashboard.html', context)


@role_required(['lecturer'])
def lecturer_student_detail(request, assignment_id):
    lecturer_profile, _ = LecturerProfile.objects.get_or_create(user=request.user)
    assignment = get_object_or_404(
        InternshipAssignment.objects.select_related(
            'application__student__user', 'application__post__department',
            'external_request__student__user', 'term',
        ),
        id=assignment_id, lecturer=lecturer_profile
    )
    weekly_reports = assignment.weekly_reports.all()
    feedback_form = LecturerFeedbackForm()

    return render(request, 'reports/lecturer_student_detail.html', {
        'assignment': assignment, 'weekly_reports': weekly_reports, 'feedback_form': feedback_form,
    })


@role_required(['lecturer'])
def lecturer_feedback_create(request, report_type, report_id):
    if request.method == 'POST':
        lecturer_profile, _ = LecturerProfile.objects.get_or_create(user=request.user)
        if report_type == 'weekly':
            report = get_object_or_404(WeeklyReport, id=report_id, assignment__lecturer=lecturer_profile)
            feedback, created = LecturerFeedback.objects.get_or_create(weekly_report=report)
        elif report_type == 'final':
            report = get_object_or_404(FinalReport, id=report_id, assignment__lecturer=lecturer_profile)
            feedback, created = LecturerFeedback.objects.get_or_create(final_report=report)
        else:
            return redirect('reports:lecturer_dashboard')

        form = LecturerFeedbackForm(request.POST, instance=feedback)
        if form.is_valid():
            form.save()
            messages.success(request, 'Lưu nhận xét thành công!')
            return redirect('reports:lecturer_student_detail', assignment_id=report.assignment.id)

    return redirect('reports:lecturer_dashboard')


# =================== STATISTICS VIEWS ===================

@role_required(['lecturer', 'admin'])
def statistics_dashboard(request):
    is_lecturer = hasattr(request.user, 'role') and request.user.role == 'lecturer'

    all_terms = InternshipTerm.objects.all().order_by('-year')
    active_term = InternshipTerm.get_active()
    departments = DepartmentProfile.objects.all()
    lecturers = LecturerProfile.objects.all()
    mentors = MentorProfile.objects.all()

    period_id = request.GET.get('period') or request.GET.get('term')
    selected_term = None
    if period_id:
        try:
            selected_term = InternshipTerm.objects.get(pk=period_id)
        except InternshipTerm.DoesNotExist:
            pass
    if not selected_term:
        selected_term = active_term or all_terms.first()

    department = request.GET.get('department', '').strip()
    lecturer_id = request.GET.get('lecturer', '')
    mentor_id = request.GET.get('mentor', '')
    pass_status = request.GET.get('pass_status', '')
    issue_status = request.GET.get('issue_status', '')

    if is_lecturer:
        try:
            lecturer_profile = LecturerProfile.objects.get(user=request.user)
            lecturer_id = str(lecturer_profile.id)
        except LecturerProfile.DoesNotExist:
            pass

    if selected_term:
        stats_data = build_statistics_context(
            period=selected_term, department=department,
            lecturer=lecturer_id, mentor=mentor_id, pass_status=pass_status,
            issue_status=issue_status,
        )
    else:
        stats_data = {
            'rows': [],
            'cards': {'total_students': 0, 'passed_count': 0, 'failed_count': 0,
                      'pass_rate': 0, 'mentor_ungraded_count': 0, 'lecturer_ungraded_count': 0,
                      'total_missing_weeks': 0, 'total_late_submissions': 0}
        }

    return render(request, 'reports/statistics.html', {
        'all_periods': all_terms, 'selected_period': selected_term,
        'departments': departments, 'lecturers': lecturers, 'mentors': mentors,
        'current_department': department, 'current_lecturer': lecturer_id,
        'current_mentor': mentor_id, 'current_pass_status': pass_status,
        'current_issue_status': issue_status,
        'is_lecturer': is_lecturer,
        'rows': stats_data['rows'], 'cards': stats_data['cards'],
    })


@role_required(['lecturer', 'admin'])
def statistics_export(request):
    is_lecturer = hasattr(request.user, 'role') and request.user.role == 'lecturer'

    all_terms = InternshipTerm.objects.all().order_by('-year')
    active_term = InternshipTerm.get_active()

    period_id = request.GET.get('period') or request.GET.get('term')
    selected_term = None
    if period_id:
        try:
            selected_term = InternshipTerm.objects.get(pk=period_id)
        except InternshipTerm.DoesNotExist:
            pass
    if not selected_term:
        selected_term = active_term or all_terms.first()

    department = request.GET.get('department', '').strip()
    lecturer_id = request.GET.get('lecturer', '')
    mentor_id = request.GET.get('mentor', '')
    pass_status = request.GET.get('pass_status', '')
    issue_status = request.GET.get('issue_status', '')

    if is_lecturer:
        try:
            lecturer_profile = LecturerProfile.objects.get(user=request.user)
            lecturer_id = str(lecturer_profile.id)
        except LecturerProfile.DoesNotExist:
            pass

    if selected_term:
        stats_data = build_statistics_context(
            period=selected_term, department=department,
            lecturer=lecturer_id, mentor=mentor_id, pass_status=pass_status,
            issue_status=issue_status,
        )
    else:
        stats_data = {'rows': []}

    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = 'attachment; filename="thong_ke_thuc_tap.csv"'
    
    writer = csv.writer(response)
    writer.writerow([
        'STT', 'Sinh viên', 'Doanh nghiệp', 'Giảng viên', 'Mentor', 
        'Điểm tuần TB', 'Final', 'Điểm tổng', 'Trạng thái', 'Số tuần thiếu', 'Số bài muộn'
    ])
    
    for idx, row in enumerate(stats_data['rows'], 1):
        student = row['student']
        student_name = student.user.get_full_name() or student.user.username if student else 'N/A'
        company = row['company'] or ''
        lecturer_name = row['lecturer'].user.get_full_name() or row['lecturer'].user.username if row.get('lecturer') else '--'
        mentor_name = row['mentor'].user.get_full_name() or row['mentor'].user.username if row.get('mentor') else '--'
        
        weekly_avg = round(row['weekly_average'], 1) if row['weekly_average'] is not None else '-'
        final_score = round(row['final_score'], 1) if row['final_score'] is not None else '-'
        total_score = round(row['total_score'], 1) if row['total_score'] is not None else '-'
        
        status = 'Đạt' if row['is_passed'] else 'Chưa đạt'
        missing_weeks = row['missing_week_count']
        late_count = row['late_submission_count']
        
        writer.writerow([
            idx, student_name, company, lecturer_name, mentor_name,
            weekly_avg, final_score, total_score, status, missing_weeks, late_count
        ])
        
    return response


@role_required(['lecturer', 'admin'])
def statistics_export_excel(request):
    is_lecturer = hasattr(request.user, 'role') and request.user.role == 'lecturer'

    all_terms = InternshipTerm.objects.all().order_by('-year')
    active_term = InternshipTerm.get_active()

    period_id = request.GET.get('period') or request.GET.get('term')
    selected_term = None
    if period_id:
        try:
            selected_term = InternshipTerm.objects.get(pk=period_id)
        except InternshipTerm.DoesNotExist:
            pass
    if not selected_term:
        selected_term = active_term or all_terms.first()

    department = request.GET.get('department', '').strip()
    lecturer_id = request.GET.get('lecturer', '')
    mentor_id = request.GET.get('mentor', '')
    pass_status = request.GET.get('pass_status', '')
    issue_status = request.GET.get('issue_status', '')

    if is_lecturer:
        try:
            lecturer_profile = LecturerProfile.objects.get(user=request.user)
            lecturer_id = str(lecturer_profile.id)
        except LecturerProfile.DoesNotExist:
            pass

    if selected_term:
        stats_data = build_statistics_context(
            period=selected_term, department=department,
            lecturer=lecturer_id, mentor=mentor_id, pass_status=pass_status,
            issue_status=issue_status,
        )
    else:
        stats_data = {
            'rows': [],
            'cards': {'total_students': 0, 'passed_count': 0, 'failed_count': 0,
                      'pass_rate': 0, 'mentor_ungraded_count': 0, 'lecturer_ungraded_count': 0,
                      'total_missing_weeks': 0, 'total_late_submissions': 0}
        }

    wb = openpyxl.Workbook()
    
    # Sheet 1: TongQuan
    ws_tongquan = wb.active
    ws_tongquan.title = "TongQuan"
    
    cards = stats_data.get('cards', {})
    
    # Write cards
    ws_tongquan.append(["THỐNG KÊ TỔNG QUAN"])
    ws_tongquan['A1'].font = Font(bold=True, size=14)
    ws_tongquan.append([])
    
    ws_tongquan.append(["Chỉ số", "Giá trị"])
    ws_tongquan['A3'].font = Font(bold=True)
    ws_tongquan['B3'].font = Font(bold=True)
    
    metrics = [
        ("Tổng sinh viên", cards.get('total_students', 0)),
        ("Số SV đạt", cards.get('passed_count', 0)),
        ("Số SV chưa đạt", cards.get('failed_count', 0)),
        ("Tỷ lệ đạt (%)", cards.get('pass_rate', 0)),
        ("Mentor chưa chấm", cards.get('mentor_ungraded_count', 0)),
        ("GV chưa chấm", cards.get('lecturer_ungraded_count', 0)),
        ("Tổng tuần thiếu", cards.get('total_missing_weeks', 0)),
        ("Tổng bài nộp muộn", cards.get('total_late_submissions', 0)),
    ]
    
    for row_idx, (label, val) in enumerate(metrics, start=4):
        ws_tongquan.cell(row=row_idx, column=1, value=label)
        ws_tongquan.cell(row=row_idx, column=2, value=val)
        
    # Chart 1: Pass/Fail (Pie)
    pie = PieChart()
    labels = Reference(ws_tongquan, min_col=1, min_row=5, max_row=6)
    data = Reference(ws_tongquan, min_col=2, min_row=4, max_row=6)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.title = "Tỷ lệ Đạt / Chưa đạt"
    ws_tongquan.add_chart(pie, "D3")
    
    # Chart 2: Issues (Bar)
    bar = BarChart()
    bar.type = "col"
    bar.style = 10
    bar.title = "Các vấn đề cần xử lý"
    bar_labels = Reference(ws_tongquan, min_col=1, min_row=8, max_row=11)
    bar_data = Reference(ws_tongquan, min_col=2, min_row=7, max_row=11)
    bar.add_data(bar_data, titles_from_data=True)
    bar.set_categories(bar_labels)
    ws_tongquan.add_chart(bar, "D18")
    
    for col in ['A', 'B']:
        ws_tongquan.column_dimensions[col].width = 25

    # Sheet 2: ChiTiet
    ws_chitiet = wb.create_sheet(title="ChiTiet")
    headers = [
        'STT', 'Sinh viên', 'Doanh nghiệp', 'Giảng viên', 'Mentor', 
        'Tuần TB', 'Final', 'Điểm tổng', 'Trạng thái', 'Tuần thiếu', 'Bài muộn', 'Mức độ rủi ro'
    ]
    ws_chitiet.append(headers)
    for cell in ws_chitiet[1]:
        cell.font = Font(bold=True)
        
    for idx, row in enumerate(stats_data['rows'], 1):
        student = row['student']
        student_name = student.user.get_full_name() or student.user.username if student else 'N/A'
        company = row['company'] or ''
        lecturer_name = row['lecturer'].user.get_full_name() or row['lecturer'].user.username if row.get('lecturer') else '--'
        mentor_name = row['mentor'].user.get_full_name() or row['mentor'].user.username if row.get('mentor') else '--'
        
        weekly_avg = round(row['weekly_average'], 1) if row['weekly_average'] is not None else '-'
        final_score = round(row['final_score'], 1) if row['final_score'] is not None else '-'
        total_score = round(row['total_score'], 1) if row['total_score'] is not None else '-'
        
        status = 'Đạt' if row['is_passed'] else 'Chưa đạt'
        missing_weeks = row['missing_week_count']
        late_count = row['late_submission_count']
        risk_label = row.get('risk_label', 'Bình thường')
        
        ws_chitiet.append([
            idx, student_name, company, lecturer_name, mentor_name,
            weekly_avg, final_score, total_score, status, missing_weeks, late_count, risk_label
        ])
        
    for col in ['B', 'C', 'D', 'E']:
        ws_chitiet.column_dimensions[col].width = 25
    for col in ['A', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
        ws_chitiet.column_dimensions[col].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"internship_statistics_{selected_term.id if selected_term else 'all'}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response


@role_required(['lecturer', 'admin'])
def statistics_assignment_detail(request, assignment_id):
    from django.core.exceptions import PermissionDenied
    
    assignment = get_object_or_404(
        InternshipAssignment.objects.select_related(
            'application__student__user', 'application__post__department',
            'external_request__student__user', 'term', 'lecturer', 'mentor_group__mentor'
        ),
        id=assignment_id
    )

    is_lecturer = hasattr(request.user, 'role') and request.user.role == 'lecturer'
    if is_lecturer:
        if not hasattr(assignment.lecturer, 'user') or assignment.lecturer.user != request.user:
            raise PermissionDenied("Bạn không có quyền xem sinh viên của giảng viên khác.")

    weekly_reports = list(assignment.weekly_reports.all().order_by('week_number'))
    for report in weekly_reports:
        report.mentor_grade = report.grades.filter(grader_type='mentor').first()
        report.lecturer_grade = report.grades.filter(grader_type='lecturer').first()

    final_report = getattr(assignment, 'final_report', None)
    if final_report:
        final_report.mentor_grade = final_report.grades.filter(grader_type='mentor').first()
        final_report.lecturer_grade = final_report.grades.filter(grader_type='lecturer').first()

    score_info = calculate_assignment_score(assignment)
    if not isinstance(score_info, dict):
        score_info = {}

    return render(request, 'reports/statistics_assignment_detail.html', {
        'assignment': assignment,
        'weekly_reports': weekly_reports,
        'final_report': final_report,
        'score_info': score_info,
        'details': score_info.get('details', {})
    })


# =================== MOODLE-STYLE DASHBOARD ===================

from .models import WeeklyDeadline, FinalDeadline, TermConfig, MentorGroup
from django.db.models import Count, Q
from django.http import JsonResponse
from django.views.decorators.http import require_POST


def _get_moodle_base_qs(request):
    """Return (assignments_qs, is_mentor, profile, selected_term, all_terms, active_term)."""
    is_mentor = request.user.role == 'mentor'
    selected_term, all_terms, active_term = get_selected_term(request)

    if is_mentor:
        profile, _ = MentorProfile.objects.get_or_create(user=request.user)
        assignments = InternshipAssignment.objects.filter(
            mentor_group__mentor=profile, is_active=True,
        )
    else:
        profile, _ = LecturerProfile.objects.get_or_create(user=request.user)
        assignments = InternshipAssignment.objects.filter(
            lecturer=profile, is_active=True,
        )

    if selected_term:
        assignments = assignments.filter(term=selected_term)

    assignments = assignments.select_related(
        'application__student__user', 'application__post__department__user',
        'external_request__student__user', 'term', 'mentor_group', 'mentor_group__mentor',
    )
    return assignments, is_mentor, profile, selected_term, all_terms, active_term


@role_required(['mentor', 'lecturer'])
def moodle_dashboard(request):
    """Trang tổng quan kiểu Moodle: grid tuần 1-10 + final + thống kê."""
    assignments, is_mentor, profile, selected_term, all_terms, active_term = _get_moodle_base_qs(request)

    total_students = assignments.count()

    # Build week summary 1-10
    weeks = []
    deadline_objs = {}
    if selected_term:
        for wd in WeeklyDeadline.objects.filter(term=selected_term):
            deadline_objs[wd.week_number] = wd

    for wn in range(1, 11):
        reports = WeeklyReport.objects.filter(
            assignment__in=assignments, week_number=wn,
        )
        submitted = reports.filter(is_submitted=True).count()
        not_submitted = total_students - submitted
        grader = 'mentor' if is_mentor else 'lecturer'
        graded = reports.filter(is_submitted=True, grades__grader_type=grader).distinct().count()
        pending = submitted - graded

        wd_obj = deadline_objs.get(wn)
        weeks.append({
            'number': wn,
            'deadline': wd_obj.deadline if wd_obj else None,
            'title': wd_obj.title if wd_obj else '',
            'description': wd_obj.description if wd_obj else '',
            'deadline_id': wd_obj.id if wd_obj else None,
            'submitted': submitted,
            'not_submitted': not_submitted,
            'graded': graded,
            'pending_grade': pending,
            'total': total_students,
        })

    # Final summary
    final_dl_obj = None
    min_weeks = 5
    if selected_term:
        try:
            final_dl_obj = FinalDeadline.objects.get(term=selected_term)
        except FinalDeadline.DoesNotExist:
            pass
        try:
            min_weeks = TermConfig.objects.get(term=selected_term).min_required_weeks
        except TermConfig.DoesNotExist:
            pass

    final_submitted = FinalReport.objects.filter(
        assignment__in=assignments, is_submitted=True,
    ).count()

    # Mentor groups for filter
    if is_mentor:
        mentor_groups = []
    else:
        mentor_groups = MentorGroup.objects.filter(
            term=selected_term, assignments__in=assignments,
        ).distinct() if selected_term else MentorGroup.objects.none()

    # === Statistics cards (similar to lecturer_dashboard) ===
    grader_type = 'mentor' if is_mentor else 'lecturer'
    all_weekly_reports = WeeklyReport.objects.filter(
        assignment__in=assignments, is_submitted=True,
    )
    total_weekly = all_weekly_reports.count()
    weekly_waiting = all_weekly_reports.exclude(grades__grader_type=grader_type).count()

    all_final_reports = FinalReport.objects.filter(
        assignment__in=assignments, is_submitted=True,
    )
    total_final = all_final_reports.count()
    final_waiting = all_final_reports.exclude(grades__grader_type=grader_type).count()

    students_no_reports = assignments.annotate(
        num_submitted=Count('weekly_reports', filter=Q(weekly_reports__is_submitted=True))
    ).filter(num_submitted=0).count()

    return render(request, 'reports/moodle/dashboard.html', {
        'weeks': weeks,
        'total_students': total_students,
        'final_deadline': final_dl_obj.deadline if final_dl_obj else None,
        'final_dl_obj': final_dl_obj,
        'final_submitted': final_submitted,
        'final_not_submitted': total_students - final_submitted,
        'min_required_weeks': min_weeks,
        'all_periods': all_terms,
        'selected_period': selected_term,
        'active_period': active_term,
        'is_mentor': is_mentor,
        'mentor_groups': mentor_groups,
        'metrics': {
            'total_students': total_students,
            'total_weekly': total_weekly,
            'weekly_waiting': weekly_waiting,
            'total_final': total_final,
            'final_waiting': final_waiting,
            'students_no_reports': students_no_reports,
            'total_waiting': weekly_waiting + final_waiting,
        },
    })


@role_required(['mentor', 'lecturer'])
def moodle_week_detail(request, week_number):
    """Chi tiết 1 tuần: bảng sinh viên + trạng thái + quick grade."""
    assignments, is_mentor, profile, selected_term, all_terms, active_term = _get_moodle_base_qs(request)

    # Filters
    tab = request.GET.get('tab', 'all')
    company_q = request.GET.get('company', '').strip()
    group_id = request.GET.get('group', '')

    if company_q:
        assignments = assignments.filter(
            Q(application__post__department__organization_name__icontains=company_q) |
            Q(application__post__department__user__username__icontains=company_q) |
            Q(external_request__company_name__icontains=company_q)
        )
    if group_id:
        assignments = assignments.filter(mentor_group_id=group_id)

    # Build rows: one per assignment
    reports_qs = WeeklyReport.objects.filter(
        assignment__in=assignments, week_number=week_number,
    ).select_related(
        'assignment__application__student__user',
        'assignment__external_request__student__user',
        'assignment__application__post__department',
        'assignment__mentor_group__mentor',
    ).prefetch_related('grades', 'comments__author', 'comments__replies__author')

    report_map = {r.assignment_id: r for r in reports_qs}

    rows = []
    for a in assignments:
        report = report_map.get(a.id)
        student = a.get_student
        mentor_grade = None
        lecturer_grade = None
        comments = []
        if report:
            mentor_grade = next((g for g in report.grades.all() if g.grader_type == 'mentor'), None)
            lecturer_grade = next((g for g in report.grades.all() if g.grader_type == 'lecturer'), None)
            comments = [c for c in report.comments.all() if c.parent_id is None]

        row = {
            'assignment': a,
            'student': student,
            'report': report,
            'is_submitted': report.is_submitted if report else False,
            'is_late': report.is_late if report else False,
            'submitted_at': report.submitted_at if report else None,
            'mentor_grade': mentor_grade,
            'lecturer_grade': lecturer_grade,
            'report_id': report.id if report else None,
            'comments': comments,
        }

        # Tab filter
        if tab == 'submitted' and not row['is_submitted']:
            continue
        if tab == 'not_submitted' and row['is_submitted']:
            continue

        rows.append(row)

    # Deadline
    deadline_obj = None
    if selected_term:
        try:
            deadline_obj = WeeklyDeadline.objects.get(term=selected_term, week_number=week_number)
        except WeeklyDeadline.DoesNotExist:
            pass

    # Mentor groups for filter (lecturer only)
    if is_mentor:
        mentor_groups = []
    else:
        mentor_groups = MentorGroup.objects.filter(
            term=selected_term,
        ).select_related('mentor') if selected_term else MentorGroup.objects.none()

    comment_form = ReportCommentForm()

    return render(request, 'reports/moodle/week_detail.html', {
        'week_number': week_number,
        'rows': rows,
        'deadline': deadline_obj.deadline if deadline_obj else None,
        'deadline_obj': deadline_obj,
        'tab': tab,
        'company_q': company_q,
        'group_id': group_id,
        'is_mentor': is_mentor,
        'all_periods': all_terms,
        'selected_period': selected_term,
        'active_period': active_term,
        'mentor_groups': mentor_groups,
        'submitted_count': sum(1 for r in rows if r['is_submitted']),
        'total_count': len(rows),
        'comment_form': comment_form,
    })


@role_required(['mentor', 'lecturer'])
def moodle_final_detail(request):
    """Chi tiết báo cáo tổng kết: bảng sinh viên + eligibility."""
    assignments, is_mentor, profile, selected_term, all_terms, active_term = _get_moodle_base_qs(request)

    tab = request.GET.get('tab', 'all')
    company_q = request.GET.get('company', '').strip()
    group_id = request.GET.get('group', '')

    if company_q:
        assignments = assignments.filter(
            Q(application__post__department__organization_name__icontains=company_q) |
            Q(external_request__company_name__icontains=company_q)
        )
    if group_id:
        assignments = assignments.filter(mentor_group_id=group_id)

    min_weeks = 5
    if selected_term:
        try:
            min_weeks = TermConfig.objects.get(term=selected_term).min_required_weeks
        except TermConfig.DoesNotExist:
            pass

    reports_qs = FinalReport.objects.filter(
        assignment__in=assignments,
    ).select_related(
        'assignment__application__student__user',
        'assignment__external_request__student__user',
        'assignment__application__post__department',
        'assignment__mentor_group__mentor',
    ).prefetch_related('grades')

    report_map = {r.assignment_id: r for r in reports_qs}

    # Prefetch weekly submitted counts
    weekly_counts = {}
    for item in WeeklyReport.objects.filter(
        assignment__in=assignments, is_submitted=True,
    ).values('assignment_id').annotate(cnt=Count('id')):
        weekly_counts[item['assignment_id']] = item['cnt']

    rows = []
    for a in assignments:
        report = report_map.get(a.id)
        student = a.get_student
        submitted_weeks = weekly_counts.get(a.id, 0)
        eligible = submitted_weeks >= min_weeks

        mentor_grade = None
        lecturer_grade = None
        if report:
            mentor_grade = next((g for g in report.grades.all() if g.grader_type == 'mentor'), None)
            lecturer_grade = next((g for g in report.grades.all() if g.grader_type == 'lecturer'), None)

        row = {
            'assignment': a,
            'student': student,
            'report': report,
            'is_submitted': report.is_submitted if report else False,
            'is_late': report.is_late if report else False,
            'submitted_at': report.submitted_at if report else None,
            'submitted_weeks': submitted_weeks,
            'eligible': eligible,
            'mentor_grade': mentor_grade,
            'lecturer_grade': lecturer_grade,
            'report_id': report.id if report else None,
        }

        if tab == 'submitted' and not row['is_submitted']:
            continue
        if tab == 'not_submitted' and row['is_submitted']:
            continue

        rows.append(row)

    deadline = None
    if selected_term:
        try:
            deadline = FinalDeadline.objects.get(term=selected_term).deadline
        except FinalDeadline.DoesNotExist:
            pass

    mentor_groups = []
    if not is_mentor and selected_term:
        mentor_groups = MentorGroup.objects.filter(term=selected_term).select_related('mentor')

    return render(request, 'reports/moodle/final_detail.html', {
        'rows': rows,
        'deadline': deadline,
        'min_required_weeks': min_weeks,
        'tab': tab,
        'company_q': company_q,
        'group_id': group_id,
        'is_mentor': is_mentor,
        'all_periods': all_terms,
        'selected_period': selected_term,
        'active_period': active_term,
        'mentor_groups': mentor_groups,
        'submitted_count': sum(1 for r in rows if r['is_submitted']),
        'total_count': len(rows),
    })


@require_POST
@role_required(['mentor', 'lecturer'])
def quick_grade(request, report_type, report_id):
    """AJAX endpoint cho quick grade từ bảng."""
    is_mentor = request.user.role == 'mentor'
    score = request.POST.get('score')
    comment = request.POST.get('comment', '')

    if not score:
        return JsonResponse({'ok': False, 'error': 'Thiếu điểm'}, status=400)

    try:
        if report_type == 'weekly':
            report = get_object_or_404(WeeklyReport, id=report_id)
            if is_mentor:
                profile, _ = MentorProfile.objects.get_or_create(user=request.user)
                grade_weekly_report_by_mentor(report, profile, score, comment)
            else:
                profile, _ = LecturerProfile.objects.get_or_create(user=request.user)
                grade_weekly_report_by_lecturer(report, profile, score, comment)
        elif report_type == 'final':
            report = get_object_or_404(FinalReport, id=report_id)
            if is_mentor:
                profile, _ = MentorProfile.objects.get_or_create(user=request.user)
                grade_final_report_by_mentor(report, profile, score, comment)
            else:
                profile, _ = LecturerProfile.objects.get_or_create(user=request.user)
                grade_final_report_by_lecturer(report, profile, score, comment)
        else:
            return JsonResponse({'ok': False, 'error': 'Invalid type'}, status=400)

        return JsonResponse({'ok': True, 'score': str(score)})
    except ValidationError as e:
        return JsonResponse({'ok': False, 'error': str(e.message)}, status=400)
    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=500)


# =================== DEADLINE EDIT VIEWS ===================

from .forms import WeeklyDeadlineForm, FinalDeadlineForm


@role_required(['lecturer', 'admin'])
def edit_weekly_deadline(request, deadline_id):
    """GV chỉnh tiêu đề, mô tả yêu cầu và deadline tuần."""
    deadline = get_object_or_404(WeeklyDeadline, id=deadline_id)

    if request.method == 'POST':
        form = WeeklyDeadlineForm(request.POST, instance=deadline)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.updated_by = request.user
            obj.save()
            messages.success(request, f'Đã cập nhật Tuần {deadline.week_number} thành công!')
            return redirect('reports:moodle_dashboard')
    else:
        form = WeeklyDeadlineForm(instance=deadline)

    return render(request, 'reports/moodle/edit_deadline.html', {
        'form': form,
        'deadline': deadline,
        'is_weekly': True,
        'week_number': deadline.week_number,
    })


@role_required(['lecturer', 'admin'])
def edit_final_deadline(request, deadline_id):
    """GV chỉnh tiêu đề, mô tả và deadline báo cáo tổng kết."""
    deadline = get_object_or_404(FinalDeadline, id=deadline_id)

    if request.method == 'POST':
        form = FinalDeadlineForm(request.POST, instance=deadline)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.updated_by = request.user
            obj.save()
            messages.success(request, 'Đã cập nhật Báo cáo Tổng kết thành công!')
            return redirect('reports:moodle_dashboard')
    else:
        form = FinalDeadlineForm(instance=deadline)

    return render(request, 'reports/moodle/edit_deadline.html', {
        'form': form,
        'deadline': deadline,
        'is_weekly': False,
    })


# =================== COMMENT VIEWS ===================

@require_POST
@role_required(['student', 'lecturer', 'mentor'])
def add_comment(request, report_type, report_id):
    """Thêm bình luận trên báo cáo tuần hoặc tổng kết."""
    form = ReportCommentForm(request.POST)
    if form.is_valid():
        comment = form.save(commit=False)
        comment.author = request.user
        if report_type == 'weekly':
            comment.weekly_report = get_object_or_404(WeeklyReport, id=report_id)
        elif report_type == 'final':
            comment.final_report = get_object_or_404(FinalReport, id=report_id)
        else:
            messages.error(request, 'Loại báo cáo không hợp lệ.')
            return redirect('reports:moodle_dashboard')
        comment.save()
        messages.success(request, 'Đã thêm bình luận thành công!')
    else:
        messages.error(request, 'Không thể thêm bình luận. Vui lòng kiểm tra lại.')

    # Redirect back to referring page
    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER')
    if next_url:
        return redirect(next_url)

    if report_type == 'weekly':
        report = WeeklyReport.objects.get(id=report_id)
        if request.user.role == 'student':
            return redirect('reports:student_week_detail', week_number=report.week_number)
        return redirect('reports:moodle_week_detail', week_number=report.week_number)
    return redirect('reports:moodle_dashboard')


@require_POST
@role_required(['student', 'lecturer', 'mentor'])
def reply_comment(request, comment_id):
    """Phản hồi bình luận."""
    parent = get_object_or_404(ReportComment, id=comment_id)
    form = ReportCommentForm(request.POST)
    if form.is_valid():
        reply = form.save(commit=False)
        reply.author = request.user
        reply.parent = parent
        reply.weekly_report = parent.weekly_report
        reply.final_report = parent.final_report
        reply.save()
        messages.success(request, 'Đã phản hồi thành công!')
    else:
        messages.error(request, 'Không thể phản hồi. Vui lòng kiểm tra lại.')

    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER')
    if next_url:
        return redirect(next_url)
    return redirect('reports:moodle_dashboard')


# =================== STUDENT MOODLE DASHBOARD ===================

@role_required(['student'])
def student_moodle_dashboard(request):
    """Dashboard dạng Moodle cho sinh viên: list tuần + final + thông tin phân công."""
    assignment = get_student_assignment(request)
    if not assignment:
        return render(request, 'reports/student_no_assignment.html')

    # Get term info
    term = assignment.term
    weekly_reports = assignment.weekly_reports.prefetch_related('grades').all()
    report_map = {r.week_number: r for r in weekly_reports}

    # Weekly deadlines
    deadline_objs = {}
    if term:
        for wd in WeeklyDeadline.objects.filter(term=term):
            deadline_objs[wd.week_number] = wd

    # Build weeks list
    weeks = []
    for wn in range(1, 11):
        report = report_map.get(wn)
        wd_obj = deadline_objs.get(wn)
        mentor_grade = None
        lecturer_grade = None
        if report:
            mentor_grade = report.grades.filter(grader_type='mentor').first()
            lecturer_grade = report.grades.filter(grader_type='lecturer').first()

        weeks.append({
            'number': wn,
            'report': report,
            'deadline': wd_obj.deadline if wd_obj else None,
            'deadline_obj': wd_obj,
            'title': wd_obj.title if wd_obj else '',
            'description': wd_obj.description if wd_obj else '',
            'is_submitted': report.is_submitted if report else False,
            'is_late': report.is_late if report else False,
            'is_locked': report.is_locked if report else False,
            'submitted_at': report.submitted_at if report else None,
            'mentor_grade': mentor_grade,
            'lecturer_grade': lecturer_grade,
        })

    # Final report
    has_final = hasattr(assignment, 'final_report')
    final_report = assignment.final_report if has_final else None
    final_dl_obj = None
    if term:
        try:
            final_dl_obj = FinalDeadline.objects.get(term=term)
        except FinalDeadline.DoesNotExist:
            pass

    final_mentor_grade = None
    final_lecturer_grade = None
    if final_report:
        final_mentor_grade = final_report.grades.filter(grader_type='mentor').first()
        final_lecturer_grade = final_report.grades.filter(grader_type='lecturer').first()

    submitted_count = sum(1 for w in weeks if w['is_submitted'])

    return render(request, 'reports/student_moodle_dashboard.html', {
        'assignment': assignment,
        'weeks': weeks,
        'submitted_count': submitted_count,
        'final_report': final_report,
        'final_dl_obj': final_dl_obj,
        'final_mentor_grade': final_mentor_grade,
        'final_lecturer_grade': final_lecturer_grade,
        'has_final': has_final,
    })


@role_required(['student'])
def student_week_detail(request, week_number):
    """Chi tiết 1 tuần của sinh viên: bài nộp + điểm + bình luận."""
    assignment = get_student_assignment(request)
    if not assignment:
        return render(request, 'reports/student_no_assignment.html')

    report = get_object_or_404(WeeklyReport, assignment=assignment, week_number=week_number)

    # Deadline info
    deadline_obj = None
    if assignment.term:
        try:
            deadline_obj = WeeklyDeadline.objects.get(term=assignment.term, week_number=week_number)
        except WeeklyDeadline.DoesNotExist:
            pass

    # Grades
    mentor_grade = report.grades.filter(grader_type='mentor').first()
    lecturer_grade = report.grades.filter(grader_type='lecturer').first()

    # Comments (top-level only, replies are prefetched)
    comments = report.comments.filter(parent__isnull=True).select_related('author').prefetch_related(
        'replies__author'
    )
    comment_form = ReportCommentForm()

    return render(request, 'reports/student_week_detail.html', {
        'assignment': assignment,
        'report': report,
        'week_number': week_number,
        'deadline_obj': deadline_obj,
        'mentor_grade': mentor_grade,
        'lecturer_grade': lecturer_grade,
        'comments': comments,
        'comment_form': comment_form,
    })


